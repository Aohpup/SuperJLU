"""Transcript import and selection rules adapted from Aohpup/JLU-score-calculator.

MIT License

Copyright (c) 2026 Li Jiaye (Aohpup)

Permission is hereby granted, free of charge, to any person obtaining a copy
of this software and associated documentation files (the "Software"), to deal
in the Software without restriction, including without limitation the rights
to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
copies of the Software, and to permit persons to whom the Software is
furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
SOFTWARE.

"""
from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
import csv
import json
import re
import unicodedata
import math
from statistics import median
from zipfile import BadZipFile

GRADE_VALUES = {"优秀": 95, "良好": 85, "中等": 75, "及格": 65, "不及格": 55}
BUILTINS = (
    {"id": "postgraduate", "name": "保研加权（不含限选课程）", "description": "剔除选修、限选、体育、通用学术英语和军事课程"},
    {"id": "scholarship", "name": "奖学金加权", "description": "必修课，剔除英语、体育、军事训练"},
    {"id": "comprehensive", "name": "所有课程", "description": "全部有学分的课程"},
)
BUILTIN_IDS = {item["id"] for item in BUILTINS}
REQUIRED_COLUMNS = ("课程名", "总成绩", "学分")
COLUMN_ALIASES = {
    "课程名": ("课程名", "课程名称", "科目", "科目名称", "课程", "教学科目", "course", "coursename", "coursetitle", "subject", "subjectname"),
    "总成绩": ("总成绩", "成绩", "分数", "得分", "最终成绩", "课程成绩", "考试成绩", "综合成绩", "总评", "总评成绩", "总分", "score", "mark", "marks", "grade", "finalscore", "finalgrade", "result"),
    "课程性质": ("课程性质", "性质", "课程类型", "课程类别", "修读类型", "修读性质", "课程属性", "类别", "coursetype", "type", "nature", "category"),
    "学分": ("学分", "课程学分", "修读学分", "学分数", "credit", "credits", "coursecredits"),
    "绩点": ("绩点", "教务绩点", "课程绩点", "成绩绩点", "gradepoint", "gradepoints", "gpa", "gp"),
    "学年学期": ("学年学期", "学期", "开课学期", "修读学期", "修读时间", "学期名称", "semester", "term", "academicterm"),
    "学年": ("学年", "学年度", "academic year", "academicyear", "year"),
}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024


@dataclass(frozen=True)
class Course:
    key: str
    term: str
    name: str
    nature: str
    grade: str
    score: float
    credit: float
    gpa: float
    score_valid: bool
    gpa_valid: bool = True

    def as_dict(self) -> dict:
        return {
            "key": self.key,
            "term": self.term,
            "name": self.name,
            "nature": self.nature,
            "grade": self.grade,
            "score": self.score,
            "credit": self.credit,
            "gpa": self.gpa,
            "score_valid": self.score_valid,
            "gpa_valid": self.gpa_valid,
            "required": nature_category(self.nature) == "必修",
            "builtin": {item["id"]: builtin_selected(item["id"], self) for item in BUILTINS},
        }


def clean_text(value: object) -> str:
    if value is None or isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def number_or_zero(value: object) -> float:
    try:
        result = float(value)
        return result if math.isfinite(result) else 0.0
    except (TypeError, ValueError):
        return 0.0


class ImportedCourses(list):
    """Course list carrying explanations for the import UI."""

    def __init__(self, notes: list[str]):
        super().__init__()
        self.import_notes = notes


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean_text(value)).lower()
    text = re.sub(r"\([^)]*\)|\[[^]]*\]", "", text)
    return re.sub(r"[\s_\-:：/]+", "", text)


def header_mapping(columns) -> dict[str, object]:
    """Match complete aliases; prefer canonical names over other aliases."""
    mapped = {}
    used = set()
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            matches = [column for column in columns if column not in used and normalize_header(column) == normalize_header(alias)]
            if len(matches) == 1:
                mapped[canonical] = matches[0]
                used.add(matches[0])
                break
            if len(matches) > 1:
                raise ValueError(f"存在多个同名的{canonical}列，请区分列名后重新导入")
    return mapped


def parse_score(value: object) -> float | None:
    text = unicodedata.normalize("NFKC", clean_text(value))
    if text in GRADE_VALUES:
        return float(GRADE_VALUES[text])
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)\s*分?", text):
        result = float(text.rstrip("分").strip())
        if 0 <= result <= 100:
            return result
    return None


def resolve_columns(columns: list[str], rows: list[list]) -> tuple[dict, list[str]]:
    if len(columns) != len(set(columns)):
        raise ValueError("成绩单有重复列名，请为每列设置不同的名称")
    mapping = header_mapping(columns)
    notes = []
    if "总成绩" not in mapping:
        candidates = []
        excluded_words = ("学号", "编号", "代码", "序号", "排名", "学时", "课时", "电话", "id", "rank", "hour", "code", "number", "平时", "期中", "期末", "考试", "补考", "重修")
        for index, column in enumerate(columns):
            if column in mapping.values() or any(word in normalize_header(column) for word in excluded_words):
                continue
            values = [clean_text(row[index]) for row in rows if index < len(row) and clean_text(row[index])]
            parsed = [parse_score(value) for value in values]
            valid = [value for value in parsed if value is not None]
            # A two-digit column alone is insufficient: require multiple rows,
            # grade-shaped contents and a single unambiguous candidate.
            if len(valid) >= 3 and len(valid) / len(values) >= .8 and median(valid) >= 50:
                candidates.append(column)
        if len(candidates) == 1:
            mapping["总成绩"] = candidates[0]
            notes.append(f"根据数据特征将“{candidates[0]}”识别为成绩列，请核对导入后的成绩。")
        elif len(candidates) > 1:
            raise ValueError("发现多个可能的成绩列：" + "、".join(map(str, candidates)) + "。请将最终成绩列改名为“成绩”或“总成绩”。")
    missing = [name for name in REQUIRED_COLUMNS if name not in mapping]
    if missing:
        raise ValueError("无法识别成绩单字段：" + "、".join(missing) + "。支持课程名称/科目、总成绩/分数/得分、学分等列名。")
    recognized = [f"{column}→{field}" for field, column in mapping.items() if field in REQUIRED_COLUMNS and str(column) != field]
    if recognized:
        notes.append("已识别字段：" + "、".join(recognized))
    if "课程性质" not in mapping:
        notes.append("未提供课程性质，显示为“未标注”；请按实际课程要求手动选择。")
    if "绩点" not in mapping:
        notes.append("未提供绩点，绩点显示为“—”；请使用本校成绩单中的绩点，不自动套用吉林大学换算。")
    return mapping, notes


def courses_from_rows(columns: list[str], rows: list[list]) -> list[Course]:
    mapping, notes = resolve_columns(columns, rows)
    positions = {name: columns.index(column) for name, column in mapping.items()}
    courses = ImportedCourses(notes)
    occurrences: dict[tuple[str, str, str], int] = {}
    for row in rows:
        def field(key: str) -> str:
            index = positions.get(key)
            return clean_text(row[index]) if index is not None and index < len(row) else ""
        name = field("课程名")
        if not name or name in {"合计", "总计", "汇总", "平均", "平均成绩"}:
            continue
        term = field("学年学期")
        year = field("学年")
        if year and year not in term:
            term = f"{year} / {term}" if term else year
        nature = field("课程性质") or "未标注"
        grade = field("总成绩")
        parsed_score = parse_score(grade)
        try:
            gpa = float(field("绩点"))
        except ValueError:
            gpa = float("nan")
        gpa_valid = math.isfinite(gpa) and gpa >= 0
        signature = (term, name, nature)
        ordinal = occurrences.get(signature, 0) + 1
        occurrences[signature] = ordinal
        key = json.dumps([term, name, nature, ordinal], ensure_ascii=False, separators=(",", ":"))
        courses.append(Course(
            key=key, term=term, name=name, nature=nature, grade=grade,
            score=parsed_score if parsed_score is not None else 0.0,
            credit=number_or_zero(field("学分")),
            gpa=float(gpa) if gpa_valid else 0.0,
            score_valid=parsed_score is not None, gpa_valid=gpa_valid,
        ))
    if not courses:
        raise ValueError("已识别表头，但没有找到课程记录")
    return courses


def courses_from_raw(raw: list[list]) -> list[Course]:
    """Find a header below school titles and student information."""
    for index in range(min(30, len(raw))):
        headers = [clean_text(value) or f"未命名列{position + 1}" for position, value in enumerate(raw[index])]
        mapping = header_mapping(headers)
        if "课程名" in mapping and "学分" in mapping:
            return courses_from_rows(headers, raw[index + 1:])
    raise ValueError("未找到成绩单表头；请确保有课程名称、成绩和学分列")


def courses_from_bytes(filename: str, data: bytes) -> list[Course]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("成绩单超过 25 MB")
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        workbook = None
        try:
            if suffix == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
                sheets = [(name, lambda name=name: [list(row) for row in workbook[name].iter_rows(values_only=True)])
                          for name in workbook.sheetnames]
            else:
                import xlrd
                try:
                    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
                except xlrd.biffh.XLRDError as error:
                    raise ValueError("无法读取 XLS 文件，请确认文件格式正确且未损坏") from error
                sheets = [(name, lambda name=name: [workbook.sheet_by_name(name).row_values(i)
                           for i in range(workbook.sheet_by_name(name).nrows)]) for name in workbook.sheet_names()]
            errors = []
            for sheet, read_rows in sheets:
                try:
                    courses = courses_from_raw(read_rows())
                except ValueError as error:
                    errors.append(f"{sheet}：{error}")
                    continue
                if len(sheets) > 1:
                    courses.import_notes.insert(0, f"已读取工作表“{sheet}”。")
                return courses
            raise ValueError("未找到可导入的成绩表。" + "；".join(errors[:3]))
        except ImportError as error:
            raise ValueError("Excel 导入需要 openpyxl / xlrd，请安装项目依赖后重试") from error
        except (BadZipFile, OSError, EOFError) as error:
            raise ValueError("无法读取 Excel 文件，请确认文件格式正确且未损坏") from error
        finally:
            if workbook is not None:
                if suffix == ".xlsx":
                    workbook.close()
                else:
                    workbook.release_resources()
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError("无法识别 CSV 文件编码")
        # Try common separators, including exported tab-separated CSV files.
        errors = []
        for delimiter in (",", ";", "\t"):
            rows = list(csv.reader(StringIO(text), delimiter=delimiter))
            try:
                return courses_from_raw(rows)
            except ValueError as error:
                errors.append(error)
        raise ValueError(str(max(errors, key=lambda error: len(str(error)))))
    raise ValueError("请选择 .xlsx、.xls 或 .csv 成绩单")


def load_courses(file_path: str | Path) -> list[Course]:
    path = Path(file_path)
    if path.stat().st_size > MAX_UPLOAD_BYTES:
        raise ValueError("成绩单超过 25 MB")
    return courses_from_bytes(path.name, path.read_bytes())


def nature_category(nature: str) -> str:
    normalized = re.sub(r"\s+", "", nature).lower()
    if normalized in {"required", "compulsory", "mandatory"} or "必修" in normalized:
        return "必修"
    if normalized in {"restricted", "restrictedelective"} or "限选" in normalized or "限定选修" in normalized or "限定性选修" in normalized:
        return "限选"
    if normalized in {"elective", "optional"} or "选修" in normalized or "任选" in normalized:
        return "选修"
    return nature


def builtin_selected(template_id: str, course: Course) -> bool:
    if course.credit <= 0:
        return False
    if template_id == "comprehensive":
        return True
    if template_id == "scholarship":
        return nature_category(course.nature) == "必修" and not any(
            word in course.name for word in ("英语", "体育", "军事训练")
        )
    if template_id == "postgraduate":
        return nature_category(course.nature) not in {"选修", "限选"} and not any(
            word in course.name for word in ("通用学术英语", "体育", "军事理论", "军事训练")
        )
    raise ValueError("未知的内置模板")


def calculate_selection(courses: list[Course], selected: list[bool]) -> dict:
    if len(courses) != len(selected):
        raise ValueError("课程和选择数量不一致")
    included = [course for course, checked in zip(courses, selected) if checked and course.credit > 0]
    credits = sum(course.credit for course in included)
    score_sum = sum(course.score * course.credit for course in included)
    gpa_sum = sum(course.gpa * course.credit for course in included)
    return {
        "count": len(included),
        "credits": credits,
        "weighted_score": score_sum / credits if credits else 0.0,
        "weighted_gpa": None if any(not course.gpa_valid for course in included) else (gpa_sum / credits if credits else 0.0),
        "score_sum": score_sum,
        "invalid_scores": sum(not course.score_valid for course in included),
    }
